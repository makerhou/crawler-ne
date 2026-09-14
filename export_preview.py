"""抓取预览导出：不连数据库，直接跑「RSS → 解码 → 抓正文」，导出 Excel 供人工核查。

用法：
    python export_preview.py                 # 默认抓 20 条
    python export_preview.py --limit 50      # 指定条数
    python export_preview.py --output x.xlsx # 指定输出文件

输出列：序号 / 标题 / 作者 / 发布时间 / URL / 抓取方式 / 正文长度 / 正文预览
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from src.article_parser import fetch_article_detail, strip_html
from src.browser_fetcher import BrowserFetcher
from src.config import Config, load_env_file
from src.logging_setup import setup_logger
from src.rss_fetcher import fetch_rss_entries
from src.scheduler import create_browser, create_nodriver, create_ua_pool
from src.url_decoder import decode_google_news_url

logger = setup_logger("reuters-crawler", "./logs", "INFO")

HEADERS = [
    "序号",
    "标题",
    "作者",
    "发布时间",
    "URL",
    "抓取方式",
    "正文长度",
    "正文预览",
]


def crawl_items(
    config: Config,
    browser: BrowserFetcher | None,
    limit: int,
    ua_pool=None,
    nodriver=None,
) -> list[dict]:
    """抓取并解析若干条目，返回结构化数据（不入库）。"""
    entries = fetch_rss_entries(
        rss_url=config.rss_url,
        timeout=config.http_timeout,
        user_agent=config.user_agent,
        proxies=config.proxies,
    )
    logger.info("RSS 拉取完成，条目数=%d，本次处理前 %d 条", len(entries), limit)

    items: list[dict] = []
    for entry in entries:
        if len(items) >= limit:
            break

        real_url = decode_google_news_url(
            entry["google_url"],
            interval=config.decode_interval,
            proxy=(config.proxies or {}).get("https") if config.proxies else None,
        )
        if not real_url:
            continue

        try:
            detail = fetch_article_detail(
                real_url,
                timeout=config.http_timeout,
                user_agent=config.user_agent,
                proxies=config.proxies,
                browser=browser,
                fetch_mode=config.fetch_mode,
                ua_pool=ua_pool,
                impersonate=config.impersonate,
                nodriver=nodriver,
            )
        except Exception as exc:
            logger.warning("抓取失败（跳过）: %s | %s", exc, real_url[:80])
            continue

        content = detail.get("content") or ""
        strategy = "trafilatura"
        if not content:
            content = strip_html(entry.get("summary"))
            strategy = "rss_fallback(降级)"

        items.append(
            {
                "title": detail.get("title") or entry.get("title") or "",
                "author": detail.get("author") or "",
                "published": entry.get("published") or detail.get("publish_time") or "",
                "url": real_url,
                "strategy": strategy,
                "length": len(content),
                "preview": content[:1000],
            }
        )
        logger.info("[%d/%d] %s", len(items), limit, (items[-1]["title"])[:60])

    return items


def write_excel(items: list[dict], output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reuters 抓取预览"

    # 表头
    ws.append(HEADERS)
    for col, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据
    for idx, item in enumerate(items, start=1):
        ws.append(
            [
                idx,
                item["title"],
                item["author"] or "-",
                item["published"] or "-",
                item["url"],
                item["strategy"],
                item["length"],
                item["preview"],
            ]
        )

    # 列宽
    widths = [6, 50, 16, 28, 60, 20, 10, 80]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 正文预览自动换行
    for row in range(2, len(items) + 2):
        ws.cell(row=row, column=8).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="抓取预览并导出 Excel")
    parser.add_argument("--limit", type=int, default=20, help="抓取条数（默认 20）")
    parser.add_argument("--output", default="", help="输出 xlsx 路径")
    parser.add_argument("--env", default=".env", help=".env 路径")
    args = parser.parse_args()

    load_env_file(args.env)
    config = Config()
    try:
        config.validate(require_db=False)
    except ValueError as exc:
        print(f"配置错误: {exc}", file=sys.stderr)
        return 2

    output = (
        Path(args.output)
        if args.output
        else Path(f"exports/reuters_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    )

    ua_pool = create_ua_pool(config)
    nodriver = create_nodriver(config, ua_pool)
    browser = create_browser(config)
    try:
        items = crawl_items(config, browser, args.limit, ua_pool, nodriver)
    finally:
        if browser is not None:
            browser.close()

    if not items:
        logger.error("未抓取到任何条目")
        return 1

    write_excel(items, output)
    stats = {
        "total": len(items),
        "full": sum(1 for i in items if i["strategy"] == "trafilatura"),
        "degraded": sum(1 for i in items if i["strategy"] != "trafilatura"),
    }
    logger.info(
        "导出完成: %s（共 %d 条，完整正文 %d 条，降级 %d 条）",
        output.resolve(),
        stats["total"],
        stats["full"],
        stats["degraded"],
    )
    print(f"\nExcel 已生成: {output.resolve()}")
    print(f"共 {stats['total']} 条 | 完整正文 {stats['full']} 条 | 降级 {stats['degraded']} 条")
    return 0


if __name__ == "__main__":
    sys.exit(main())
